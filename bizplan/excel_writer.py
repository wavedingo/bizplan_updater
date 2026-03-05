import re
import shutil
from pathlib import Path
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

# These sheets must never be written to
LOCKED_SHEETS = frozenset({
    "Summary Financials",
    "Income Statement",
    "Balance Sheet",
    "Cash Flow",
})

# Map abbreviated month headers to YYYY-MM
_MONTH_ABBR = {
    "jan": "01", "feb": "02", "mar": "03", "apr": "04",
    "may": "05", "jun": "06", "jul": "07", "aug": "08",
    "sep": "09", "oct": "10", "nov": "11", "dec": "12",
}

def _header_to_yyyymm(header: str) -> str | None:
    """Convert 'Jan 2026' or 'January 2026' to '2026-01', or None if unparseable."""
    m = re.match(r"([A-Za-z]+)\s+(\d{4})", str(header).strip())
    if not m:
        return None
    month_abbr = m.group(1)[:3].lower()
    year = m.group(2)
    month_num = _MONTH_ABBR.get(month_abbr)
    if not month_num:
        return None
    return f"{year}-{month_num}"

def find_month_column(ws: Worksheet, month: str, header_row: int = 1) -> int | None:
    """Return the 1-based column index for the given YYYY-MM month, or None."""
    for col in range(1, ws.max_column + 1):
        cell_val = ws.cell(row=header_row, column=col).value
        if cell_val and _header_to_yyyymm(str(cell_val)) == month:
            return col
    return None

def find_row_by_label(ws: Worksheet, label: str, label_col: int = 1) -> int | None:
    """Return the 1-based row index where column label_col matches label, or None."""
    for row in range(1, ws.max_row + 1):
        cell_val = ws.cell(row=row, column=label_col).value
        if cell_val and str(cell_val).strip() == label.strip():
            return row
    return None

def scan_workbook_structure(wb_path: Path) -> dict:
    """Scan writable input sheets and return their header_row and label_col positions.

    Returns: {sheet_name: {"header_row": int, "label_col": int}}
    Heuristic: header_row is the first row containing a recognizable month string.
    """
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    structure = {}
    for sheet_name in wb.sheetnames:
        if sheet_name in LOCKED_SHEETS:
            continue
        ws = wb[sheet_name]
        header_row = None
        for row_idx in range(1, min(10, ws.max_row + 1)):
            for col_idx in range(1, ws.max_column + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val and _header_to_yyyymm(str(val)):
                    header_row = row_idx
                    break
            if header_row:
                break
        if header_row:
            structure[sheet_name] = {"header_row": header_row, "label_col": 1}
    wb.close()
    return structure

def write_cell_value(wb_path: Path, sheet_name: str, row: int, col: int,
                     value: float, output_path: Path) -> None:
    """Copy workbook to output_path and write value to sheet_name[row][col].

    Raises ValueError if sheet_name is locked.
    """
    if sheet_name in LOCKED_SHEETS:
        raise ValueError(f"Sheet '{sheet_name}' is locked and cannot be written to.")
    shutil.copy2(wb_path, output_path)
    wb = openpyxl.load_workbook(output_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in workbook.")
    wb[sheet_name].cell(row=row, column=col, value=value)
    wb.save(output_path)

def write_monthly_actuals(
    source_wb_path: Path,
    output_wb_path: Path,
    actuals: list[dict],
    model_structure: dict,
    month: str,
) -> int:
    """Write all actuals for a month to a versioned workbook copy.

    actuals: list of {excel_sheet, excel_row_label, amount}
    model_structure: from scan_workbook_structure()
    Returns: number of cells written.
    """
    if source_wb_path != output_wb_path:
        shutil.copy2(source_wb_path, output_wb_path)

    wb = openpyxl.load_workbook(output_wb_path)
    cells_written = 0

    for item in actuals:
        sheet_name = item["excel_sheet"]
        row_label = item["excel_row_label"]
        amount = item["amount"]

        if sheet_name not in wb.sheetnames:
            continue
        if sheet_name in LOCKED_SHEETS:
            continue
        if sheet_name not in model_structure:
            continue

        ws = wb[sheet_name]
        sheet_meta = model_structure[sheet_name]
        header_row = sheet_meta["header_row"]
        label_col = sheet_meta["label_col"]

        col = find_month_column(ws, month, header_row=header_row)
        row = find_row_by_label(ws, row_label, label_col=label_col)

        if col is None or row is None:
            continue

        ws.cell(row=row, column=col, value=abs(amount))  # Store as positive in model
        cells_written += 1

    wb.save(output_wb_path)
    return cells_written

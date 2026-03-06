"""End-to-end test using the real sample QuickBooks CSV and a minimal test workbook."""
import pytest
import openpyxl
import shutil
from pathlib import Path
from click.testing import CliRunner
from bizplan.cli import cli

SAMPLE_CSV = Path("quickbooks_exports/TD_Bank-3.csv")


@pytest.fixture
def env(tmp_path):
    """Set up a complete test environment with a sample workbook and QB CSV."""
    # Create minimal workbook matching known QB categories
    wb = openpyxl.Workbook()
    # Marketing Exp sheet
    ws1 = wb.active
    ws1.title = "Marketing Exp"
    ws1["A1"] = "Category"
    ws1["B1"] = "Jan 2026"
    ws1["C1"] = "Feb 2026"
    ws1["A2"] = "Social Media"
    ws1["A3"] = "Website Ads"
    # Production Exp sheet
    ws2 = wb.create_sheet("Production Exp")
    ws2["A1"] = "Category"
    ws2["B1"] = "Jan 2026"
    ws2["A2"] = "Software & Tools"
    ws2["A3"] = "Wardrobe"

    wb_path = tmp_path / "model.xlsx"
    wb.save(wb_path)

    # Copy QB CSV
    qb_dir = tmp_path / "quickbooks_exports"
    qb_dir.mkdir()
    if SAMPLE_CSV.exists():
        shutil.copy(SAMPLE_CSV, qb_dir / "TD_Bank-3.csv")

    return {"tmp_path": tmp_path, "wb_path": wb_path}


@pytest.mark.skipif(not SAMPLE_CSV.exists(), reason="Sample CSV not present")
def test_full_pipeline(env, monkeypatch):
    monkeypatch.chdir(env["tmp_path"])
    runner = CliRunner()

    # Init
    result = runner.invoke(cli, ["init", "--workbook", str(env["wb_path"]), "--force"])
    assert result.exit_code == 0, result.output

    # Update (no-prompt so we skip interactive categorization)
    result = runner.invoke(
        cli,
        ["update", "--month", "2026-01", "--workbook", str(env["wb_path"]), "--no-prompt"],
    )
    assert result.exit_code == 0, result.output
    assert "transactions" in result.output.lower()

    # Output workbook created
    assert (env["tmp_path"] / "output" / "workbooks" / "2026-01_update.xlsx").exists()

    # Report created
    report_path = env["tmp_path"] / "output" / "reports" / "2026-01-report.md"
    assert report_path.exists()
    content = report_path.read_text()
    assert "2026-01" in content

    # Deduplication: re-run should produce same totals (no new transactions)
    result2 = runner.invoke(
        cli,
        ["update", "--month", "2026-01", "--workbook", str(env["wb_path"]), "--no-prompt"],
    )
    assert result2.exit_code == 0
    # Second run: 0 new transactions (all duplicates)
    assert "0 new" in result2.output or "duplicates skipped" in result2.output

    # Output workbook opens cleanly
    wb = openpyxl.load_workbook(env["tmp_path"] / "output" / "workbooks" / "2026-01_update.xlsx")
    assert wb is not None
    wb.close()


@pytest.mark.skipif(not SAMPLE_CSV.exists(), reason="Sample CSV not present")
def test_map_list_after_update(env, monkeypatch):
    """After an update run, map --list should show confirmed mappings (if any)."""
    monkeypatch.chdir(env["tmp_path"])
    runner = CliRunner()
    runner.invoke(cli, ["init", "--workbook", str(env["wb_path"]), "--force"])
    runner.invoke(
        cli,
        ["update", "--month", "2026-01", "--workbook", str(env["wb_path"]), "--no-prompt"],
    )
    result = runner.invoke(cli, ["map", "--list"])
    assert result.exit_code == 0

import pytest
import openpyxl
from pathlib import Path
from click.testing import CliRunner
from bizplan.cli import cli

@pytest.fixture
def sample_wb(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Marketing Exp"
    ws["A1"] = "Category"
    ws["B1"] = "Jan 2026"
    ws["A2"] = "Social Media"
    # Add a locked sheet
    wb.create_sheet("Income Statement")
    path = tmp_path / "model.xlsx"
    wb.save(path)
    return path

def test_init_creates_db_and_config(tmp_path, sample_wb, monkeypatch):
    monkeypatch.chdir(tmp_path)
    runner = CliRunner()
    result = runner.invoke(cli, ["init", "--workbook", str(sample_wb), "--force"])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "data" / "transactions.db").exists()
    assert (tmp_path / "config" / "model.yaml").exists()

def test_validate_passes_clean_model(tmp_path, sample_wb, monkeypatch):
    monkeypatch.chdir(tmp_path)
    runner = CliRunner()
    runner.invoke(cli, ["init", "--workbook", str(sample_wb), "--force"])
    result = runner.invoke(cli, ["validate", "--workbook", str(sample_wb)])
    assert result.exit_code == 0

def test_help_flag(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    runner = CliRunner()
    result = runner.invoke(cli, ["--help"])
    assert "init" in result.output
    assert result.exit_code == 0
